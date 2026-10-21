from __future__ import annotations

from pathlib import Path

from src.rag.parsers.base import Block, DocumentParser, register
from src.utils.logger import logger


@register
class XlsxParser(DocumentParser):
    supported = (".xlsx", ".xlsm")

    def parse(self, path: Path) -> list:
        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.warning("缺少 openpyxl 依赖，跳过 .xlsx 解析")
            return []

        wb = load_workbook(str(path), read_only=True, data_only=True)
        blocks: list = []
        try:
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                header = [str(c) if c is not None else "" for c in rows[0]]
                for r_idx, row in enumerate(rows[1:], start=2):
                    cells = [str(c) if c is not None else "" for c in row]
                    if not any(c.strip() for c in cells):
                        continue
                    # 列名前置，避免单元格值脱离表头失去语义
                    record = "; ".join(
                        f"{header[j]}: {cells[j]}"
                        for j in range(min(len(header), len(cells)))
                        if cells[j].strip()
                    )
                    if record:
                        blocks.append(Block(type="paragraph", text=record, meta={"sheet": ws.title, "row": r_idx}))
        finally:
            wb.close()
        return blocks
