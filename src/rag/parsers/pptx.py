from __future__ import annotations

from pathlib import Path

from src.rag.parsers.base import Block, DocumentParser, register
from src.utils.logger import logger


@register
class PptxParser(DocumentParser):
    supported = (".pptx",)

    def parse(self, path: Path) -> list:
        try:
            from pptx import Presentation
        except ImportError:
            logger.warning("缺少 python-pptx 依赖，跳过 .pptx 解析")
            return []

        prs = Presentation(str(path))
        blocks: list = []
        for idx, slide in enumerate(prs.slides, 1):
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame and title_shape.text.strip():
                blocks.append(Block(type="heading", text=title_shape.text.strip(), level=1, meta={"slide": idx}))
            for shape in slide.shapes:
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if shape.has_text_frame and shape.text.strip():
                    blocks.append(Block(type="paragraph", text=shape.text.strip(), meta={"slide": idx}))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    blocks.append(Block(type="paragraph", text=notes, meta={"slide": idx, "is_notes": True}))
        return blocks
